"""P4 测试：定时计划、Word 导出、多语言。"""

from __future__ import annotations

import io
import sys
from datetime import datetime, timedelta
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "apps"))

from common import exporters, generation, i18n, store


@pytest.fixture()
def sandbox(tmp_path, monkeypatch):
    data_dir = tmp_path / ".stride-gpt"
    monkeypatch.setattr(store, "DATA_DIR", data_dir)
    monkeypatch.setattr(store, "CONFIG_FILE", data_dir / "web_config.json")
    monkeypatch.setattr(store, "KEY_FILE", data_dir / "web_key")
    monkeypatch.setattr(store, "DB_FILE", data_dir / "web.db")
    return data_dir


# ---------------------------------------------------------------------------
# 定时计划
# ---------------------------------------------------------------------------


class TestSchedules:
    def test_crud_and_due(self, sandbox):
        sid = store.add_schedule("项目X", 60, True)
        schedule = store.get_schedule(sid)
        assert schedule["project"] == "项目X"
        assert schedule["interval_minutes"] == 60
        # 从未运行过 → 立即到期
        assert [s["id"] for s in store.due_schedules()] == [sid]

        store.update_schedule(sid, last_run=store._now(), last_status="ok")
        assert store.due_schedules() == []  # 刚跑过，未到期

        past = (datetime.now() - timedelta(minutes=61)).strftime("%Y-%m-%d %H:%M:%S")
        store.update_schedule(sid, last_run=past)
        assert [s["id"] for s in store.due_schedules()] == [sid]

        store.update_schedule(sid, enabled=0)
        assert store.due_schedules() == []  # 暂停后不再到期

        store.delete_schedule(sid)
        assert store.list_schedules() == []

    def test_add_is_upsert(self, sandbox):
        first = store.add_schedule("项目Y", 30)
        second = store.add_schedule("项目Y", 120, False)
        assert first == second
        assert store.get_schedule(first)["interval_minutes"] == 120


class TestScheduledRun:
    def test_run_updates_status_and_notifies(self, sandbox, monkeypatch):
        store.upsert_project("项目Z", "一个订单系统")
        sid = store.add_schedule("项目Z", 60, True)

        calls = []

        def fake_pipeline(**kwargs):
            calls.append(kwargs["project"])
            generation.archive_result(
                "threat_model",
                "# tm",
                {"threat_model": [{"Threat Type": "Tampering", "Scenario": "旧威胁"}]},
                project=kwargs["project"],
            )
            generation.archive_result(
                "threat_model",
                "# tm",
                {"threat_model": [{"Threat Type": "Spoofing", "Scenario": "新威胁"}]},
                project=kwargs["project"],
            )
            return [{"kind": "threat_model", "ok": True}]

        pushed = []
        monkeypatch.setattr(generation, "run_full_pipeline", fake_pipeline)
        monkeypatch.setattr(generation, "push_text", lambda title, text, webhooks=None: pushed.append(text) or [{"ok": True}])

        result = generation.run_scheduled_analysis(store.get_schedule(sid))
        assert calls == ["项目Z"]
        assert result["status"] in ("ok", "partial")
        assert result["pushed"] == 1
        assert any("新威胁" in text for text in pushed)
        updated = store.get_schedule(sid)
        assert updated["last_run"]
        assert "ok" in updated["last_status"] or "partial" in updated["last_status"]

    def test_missing_description_skips(self, sandbox):
        store.upsert_project("空项目", "")
        sid = store.add_schedule("空项目", 60)
        result = generation.run_scheduled_analysis(store.get_schedule(sid))
        assert result["status"] == "skipped"
        assert "缺少应用描述" in store.get_schedule(sid)["last_status"]


# ---------------------------------------------------------------------------
# Word 导出
# ---------------------------------------------------------------------------


class TestDocxExport:
    def test_docx_contains_title_and_tables(self, sandbox):
        from docx import Document

        report = {
            "id": 1,
            "kind": "threat_model",
            "project": "演示项目",
            "created_at": "2026-10-16 12:00:00",
            "provider": "OpenAI API",
            "model": "gpt-6-astra",
            "content_md": (
                "# 威胁模型\n\n"
                "| 威胁类型 | 攻击场景 |\n|---|---|\n| Spoofing | 伪造令牌 |\n\n"
                "- 建议一\n\n```\ncode block\n```\n"
            ),
        }
        data = exporters.markdown_to_docx(report)
        assert data[:2] == b"PK"  # docx 是 zip 容器

        doc = Document(io.BytesIO(data))
        texts = [p.text for p in doc.paragraphs]
        assert any("演示项目" in t for t in texts)
        assert any(t == "威胁模型" for t in texts)
        table_texts = [
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        ]
        assert "伪造令牌" in table_texts

    def test_template_brand_and_footer(self, sandbox):
        from docx import Document

        store.save_config({"report_brand": "ACME 安全部", "report_footer": "机密 · 禁止外传"})
        report = {
            "id": 2,
            "kind": "dfd",
            "project": "P",
            "created_at": "t",
            "provider": "",
            "model": "",
            "content_md": "hello",
        }
        doc = Document(io.BytesIO(exporters.markdown_to_docx(report)))
        assert any("ACME 安全部" in p.text for p in doc.paragraphs)
        footer = doc.sections[0].footer.paragraphs[0].text
        assert "机密 · 禁止外传" in footer

    def test_filename_is_ascii_safe(self, sandbox):
        name = exporters.suggested_filename({"id": 9, "kind": "dfd", "project": "演示/项目"}, "pdf")
        assert name.endswith(".pdf")
        assert "/" not in name


# ---------------------------------------------------------------------------
# 多语言
# ---------------------------------------------------------------------------


class TestI18n:
    def test_switch_language(self):
        i18n.set_language("zh")
        assert i18n.t("中文", "English") == "中文"
        i18n.set_language("en")
        assert i18n.t("中文", "English") == "English"
        i18n.set_language("unknown")
        assert i18n.current_language() == "zh"  # 未知语言回退
        assert set(i18n.LANGUAGES) == {"zh", "en"}


# ---------------------------------------------------------------------------
# 通知文本
# ---------------------------------------------------------------------------


class TestNotificationPayloads:
    def test_text_payload_shapes(self):
        feishu = store.text_webhook_payload("feishu", "标题", "正文")
        assert feishu["msg_type"] == "text" and "正文" in feishu["content"]["text"]
        wecom = store.text_webhook_payload("wecom", "标题", "正文")
        assert wecom["msgtype"] == "text"
        generic = store.text_webhook_payload("generic", "标题", "正文")
        assert generic["event"] == "notification" and generic["title"] == "标题"
