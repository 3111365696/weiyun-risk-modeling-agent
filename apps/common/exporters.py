"""报告导出：Markdown → Word（python-docx）与 HTML → PDF（Chrome 无头打印）。

模板：标题页 + 元信息表 + 正文（标题/表格/列表/代码块）+ 页脚（保密标识 + 页码）。
"""

from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

from . import store

BRAND_FALLBACK = "威云-风险建模智能体"
FOOTER_FALLBACK = "内部资料 · 注意保密"

_CHROME_CANDIDATES = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "google-chrome",
    "google-chrome-stable",
    "chromium",
    "chromium-browser",
    "microsoft-edge",
]


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------


def _set_cjk_font(style, name: str = "PingFang SC", fallback: str = "Microsoft YaHei") -> None:
    from docx.oxml.ns import qn

    style.font.name = fallback
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style.element.rPr.rFonts.set(qn("w:ascii"), fallback)
    style.element.rPr.rFonts.set(qn("w:hAnsi"), fallback)


def _add_page_number(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def _shade(cell, color: str = "EEEAF8") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def markdown_to_docx(
    report: dict,
    *,
    brand: str | None = None,
    footer: str | None = None,
) -> bytes:
    """把一条报告渲染为 Word 文档（返回 docx 字节）。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    saved = store.load_config()
    brand = brand or saved.get("report_brand") or BRAND_FALLBACK
    footer = footer or saved.get("report_footer") or FOOTER_FALLBACK

    doc = Document()
    _set_cjk_font(doc.styles["Normal"], name="PingFang SC")

    # 标题页
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(brand)
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x6C, 0x59, 0xC9)

    kind_label = store.KIND_LABELS.get(report.get("kind", ""), report.get("kind", ""))
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{kind_label} · {report.get('project') or '未命名'}")
    run.font.size = Pt(24)
    run.bold = True

    meta = [
        ("项目", report.get("project") or "未命名"),
        ("报告类型", kind_label),
        ("生成时间", report.get("created_at") or ""),
        ("模型", f"{report.get('provider') or ''} / {report.get('model') or ''}".strip(" /")),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for key, value in meta:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = str(value)
    doc.add_paragraph()

    # 正文：逐行解析常见 markdown 结构
    lines = (report.get("content_md") or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            para = doc.add_paragraph()
            run = para.add_run("\n".join(block))
            run.font.name = "Menlo"
            run.font.size = Pt(9)
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            header, *body = rows[1:] if len(rows) > 1 else [rows[0]]
            cols = max(len(r) for r in rows)
            t = doc.add_table(rows=0, cols=cols)
            t.style = "Light Grid Accent 1"
            hdr = t.add_row().cells
            for idx, cell_text in enumerate(header + [""] * (cols - len(header))):
                hdr[idx].text = cell_text
                _shade(hdr[idx])
            for row in body:
                cells = t.add_row().cells
                for idx, cell_text in enumerate(row + [""] * (cols - len(row))):
                    cells[idx].text = cell_text
            doc.add_paragraph()
            continue

        if stripped.startswith(("- ", "* ")):
            doc.add_paragraph(stripped[2:], style="List Bullet")
            i += 1
            continue

        if stripped.startswith(">"):
            para = doc.add_paragraph(stripped.lstrip("> ").strip(), style="Intense Quote")
            i += 1
            continue

        for prefix, level in (("#### ", 4), ("### ", 3), ("## ", 2), ("# ", 1)):
            if stripped.startswith(prefix):
                doc.add_heading(stripped[len(prefix) :], level=level)
                break
        else:
            doc.add_paragraph(_strip_inline(stripped))
        i += 1

    # 页脚：保密标识 + 页码
    section = doc.sections[0]
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_para.text = f"{footer} · {brand} · 第 "
    _add_page_number(footer_para)
    footer_para.add_run(" 页")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _strip_inline(text: str) -> str:
    """去掉行内 markdown 标记（**bold**、`code`、链接等）。"""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return re.sub(r"\[(.+?)\]\((.+?)\)", r"\1", text)


# ---------------------------------------------------------------------------
# PDF（Chrome 无头打印）
# ---------------------------------------------------------------------------


def find_chrome() -> str | None:
    for candidate in _CHROME_CANDIDATES:
        if Path(candidate).exists():
            return candidate
        found = shutil.which(candidate)
        if found:
            return found
    return None


def html_to_pdf(html_text: str, *, timeout: int = 90) -> bytes:
    """用 Chrome/Chromium 无头模式把 HTML 打印为 PDF；未安装时抛 RuntimeError。"""
    chrome = find_chrome()
    if not chrome:
        raise RuntimeError("未找到 Chrome/Chromium，无法生成 PDF（可安装 Chrome 或改用 HTML/Word 导出）")

    with tempfile.TemporaryDirectory() as tmp:
        html_path = Path(tmp) / "report.html"
        pdf_path = Path(tmp) / "report.pdf"
        html_path.write_text(html_text, encoding="utf-8")
        cmd = [
            chrome,
            "--headless=new",
            "--disable-gpu",
            "--no-pdf-header-footer",
            "--no-first-run",
            f"--user-data-dir={tmp}/profile",
            f"--print-to-pdf={pdf_path}",
            html_path.as_uri(),
        ]
        proc = subprocess.run(cmd, capture_output=True, timeout=timeout)
        if not pdf_path.exists():
            raise RuntimeError(
                f"PDF 生成失败：{(proc.stderr or b'').decode('utf-8', 'ignore')[:200]}"
            )
        return pdf_path.read_bytes()


def report_to_pdf(report: dict, *, brand: str | None = None) -> bytes:
    return html_to_pdf(store.report_to_html(report))


def suggested_filename(report: dict, suffix: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d")
    project = (report.get("project") or "report").replace("/", "_")[:24]
    return f"{project}-{report.get('kind', 'report')}-{report.get('id', 0)}-{stamp}.{suffix}"
